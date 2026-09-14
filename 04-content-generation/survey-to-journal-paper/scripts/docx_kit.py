# -*- coding: utf-8 -*-
"""中文期刊 DOCX 排版工具箱（可复用）
用法：from docx_kit import set_font, add_three_line_table, fix_quotes_cn, setup_doc
依赖：python-docx
"""
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SONG, HEI, KAI, TNR = '宋体', '黑体', '楷体', 'Times New Roman'


# ---------- 字体 ----------
def set_font(run, cn=SONG, en=TNR, size=12, bold=False):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rPr.append(rf)
    rf.set(qn('w:ascii'), en); rf.set(qn('w:hAnsi'), en); rf.set(qn('w:eastAsia'), cn)


def setup_doc(doc, top=2.4, bottom=2.4, left=2.7, right=2.7):
    s = doc.sections[0]
    s.top_margin, s.bottom_margin = Cm(top), Cm(bottom)
    s.left_margin, s.right_margin = Cm(left), Cm(right)
    st = doc.styles['Normal']
    st.font.name = TNR
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), SONG)


# ---------- 三线表 ----------
# 三线表规范（中文社科期刊）：
#   顶线 1.5pt（w:sz=12）、表头下线 0.75pt（w:sz=6）、底线 1.5pt（w:sz=12）
#   其余全部无边框
# 稳健做法：顶线/底线画在「表级 tblBorders」，表头下线画在「表头行单元格 tcBorders」。
# 这样即使做纵向合并（vMerge）也不会断线——这是单元格级实现最常见的坑。

TRIPLE_TOP = 12      # 1.5pt 顶线
TRIPLE_MID = 6       # 0.75pt 表头下线
TRIPLE_BOTTOM = 12   # 1.5pt 底线


def clear_table_borders(table):
    """清空表级边框（三线表只需要 top/bottom，其余 none）。"""
    tblPr = table._tbl.tblPr
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:%s' % e)
        el.set(qn('w:val'), 'none'); el.set(qn('w:sz'), '0')
        el.set(qn('w:space'), '0'); el.set(qn('w:color'), 'auto')
        b.append(el)
    tblPr.append(b)


def set_table_edge(table, edge, sz):
    """在「表级」设置某条边（top/bottom）为实线。"""
    tblPr = table._tbl.tblPr
    b = tblPr.find(qn('w:tblBorders'))
    if b is None:
        b = OxmlElement('w:tblBorders'); tblPr.append(b)
    tag = qn('w:%s' % edge)
    el = b.find(tag)
    if el is None:
        el = OxmlElement('w:%s' % edge); b.append(el)
    el.set(qn('w:val'), 'single')
    el.set(qn('w:sz'), str(sz))
    el.set(qn('w:space'), '0')
    el.set(qn('w:color'), '000000')


def set_cell_borders(cell, top=None, bottom=None,
                     left=None, right=None, insideH=None, insideV=None):
    """单元格级边框。传 sz 值即画实线，None 表示不动。"""
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders')
        # tcBorders 必须排在 tcW 之后（schema 顺序），否则 Word 可能忽略
        tcW = tcPr.find(qn('w:tcW'))
        if tcW is not None:
            tcW.addnext(b)
        else:
            tcPr.append(b)
    for edge, spec in (('top', top), ('bottom', bottom),
                       ('left', left), ('right', right),
                       ('insideH', insideH), ('insideV', insideV)):
        if spec is None:
            continue
        tag = qn('w:%s' % edge)
        el = b.find(tag)
        if el is None:
            el = OxmlElement('w:%s' % edge); b.append(el)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(spec))
        el.set(qn('w:space'), '0'); el.set(qn('w:color'), '000000')


def repeat_as_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:tblHeader'); e.set(qn('w:val'), 'true')
    trPr.append(e)


def set_cant_split(row):
    """禁止该行跨页被拦腰截断（三线表数据行必设）。"""
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn('w:cantSplit')) is None:
        trPr.append(OxmlElement('w:cantSplit'))


def keep_table_together(table, on=True):
    """尽量让整张表不被分页拆开：
    给表内所有段落设 keep_with_next，最后一行的段落除外。
    """
    if not on:
        return
    rows = table.rows
    for ri, row in enumerate(rows):
        last = (ri == len(rows) - 1)
        for c in row.cells:
            for p in c.paragraphs:
                pf = p.paragraph_format
                pf.keep_with_next = not last


def set_row_height(row, cm, exact=False):
    """设置行高（exact=True 为固定值，False 为最小值）。"""
    trPr = row._tr.get_or_add_trPr()
    h = trPr.find(qn('w:trHeight'))
    if h is None:
        h = OxmlElement('w:trHeight'); trPr.append(h)
    h.set(qn('w:val'), str(int(cm * 567)))
    h.set(qn('w:hRule'), 'exact' if exact else 'atLeast')


def _sync_grid(table, widths):
    """让 tblGrid / 各行 tcW 一致（避免 Word 自动重排导致列宽乱跳）。"""
    total_twips = int(sum(widths) * 567)  # 1cm = 567 twips
    # tblW 固定为总宽
    tblPr = table._tbl.tblPr
    for tag in ('w:tblW',):
        el = tblPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); tblPr.append(el)
        el.set(qn('w:type'), 'dxa'); el.set(qn('w:w'), str(total_twips))
    # tblGrid
    grid = table._tbl.find(qn('w:tblGrid'))
    if grid is not None:
        table._tbl.remove(grid)
    grid = OxmlElement('w:tblGrid')
    for w in widths:
        gc = OxmlElement('w:gridCol'); gc.set(qn('w:w'), str(int(w * 567)))
        grid.append(gc)
    # tblGrid 必须紧跟 tblPr
    tblPr.addnext(grid)
    # 每格 tcW
    for row in table.rows:
        seen = set()
        for i, c in enumerate(row.cells):
            if id(c._tc) in seen:
                continue
            seen.add(id(c._tc))
            tcPr = c._tc.get_or_add_tcPr()
            el = tcPr.find(qn('w:tcW'))
            if el is None:
                el = OxmlElement('w:tcW')
                tcPr.insert(0, el)
            span = 1
            gs = tcPr.find(qn('w:gridSpan'))
            if gs is not None:
                span = int(gs.get(qn('w:val')))
            el.set(qn('w:type'), 'dxa')
            el.set(qn('w:w'), str(int(sum(widths[i:i + span]) * 567)))


def add_three_line_table(doc, header, rows, widths=None, size=9,
                         merges=None, first_col_left=False,
                         left_cols=(), align_map=None,
                         row_gap=1.2, keep_together=True,
                         vcenter=True):
    """标准三线表。

    参数
    ----
    widths     : 各列宽度（cm）。给了宽度就同步 tblGrid + 每格 tcW。
    merges     : [(r1, r2, col)] 以【数据行】索引（0 起）做纵向合并。
    first_col_left : 首列左对齐（数据类三线表常用）。默认 False（居中）。
    left_cols  : 需要左对齐的列索引集合（覆盖 first_col_left），如 (0, 1)。
    align_map  : {列号: 'left'/'center'/'right'} 精细控制。
    row_gap    : 单元格上下留白（pt）。表很高时调小（如 0.6）以便一页放下。
    keep_together : 尽量不让表格被分页拆开（表很高时可设 False）。
    vcenter    : 单元格内容垂直居中（合并单元格尤其需要）。

    生成后结构
    ----------
    顶线（表级 1.5pt）→ 表头行 → 表头下线（单元格 0.75pt）→ 数据行 → 底线（表级 1.5pt）
    """
    ncol = len(header)
    t = doc.add_table(rows=1, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    clear_table_borders(t)

    # 顶线 + 底线：画在表级，合并也不断
    set_table_edge(t, 'top', TRIPLE_TOP)
    set_table_edge(t, 'bottom', TRIPLE_BOTTOM)

    def _align_of(i):
        if align_map and i in align_map:
            return {'left': WD_ALIGN_PARAGRAPH.LEFT,
                    'center': WD_ALIGN_PARAGRAPH.CENTER,
                    'right': WD_ALIGN_PARAGRAPH.RIGHT}[align_map[i]]
        if i in left_cols or (i == 0 and first_col_left):
            return WD_ALIGN_PARAGRAPH.LEFT
        return WD_ALIGN_PARAGRAPH.CENTER

    def _vcenter(c):
        if not vcenter:
            return
        tcPr = c._tc.get_or_add_tcPr()
        if tcPr.find(qn('w:vAlign')) is None:
            el = OxmlElement('w:vAlign'); el.set(qn('w:val'), 'center')
            tcPr.append(el)

    # 表头行：黑体加粗，底部 0.75pt
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.space_before = Pt(2); pf.space_after = Pt(2); pf.line_spacing = 1.0
        set_font(p.add_run(str(h)), HEI, TNR, size, True)
        set_cell_borders(c, bottom=TRIPLE_MID)
        _vcenter(c)
    repeat_as_header(t.rows[0])
    set_cant_split(t.rows[0])

    # 数据行：宋体，无边框
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            c = cells[i]
            p = c.paragraphs[0]
            p.alignment = _align_of(i)
            pf = p.paragraph_format
            pf.space_before = Pt(row_gap); pf.space_after = Pt(row_gap)
            pf.line_spacing = 1.0
            set_font(p.add_run('' if v is None else str(v)), SONG, TNR, size, False)
            _vcenter(c)
        set_cant_split(t.rows[-1])

    # 纵向合并（数据行索引 +1 偏移表头）
    if merges:
        for r1, r2, col in merges:
            a = t.cell(r1 + 1, col)
            b = t.cell(r2 + 1, col)
            if a is not b:
                a.merge(b)
                _vcenter(a)

    # 逐行重排：确保「首行表头下线」没被 set_cell_borders 覆盖丢失
    for i, c in enumerate(t.rows[0].cells):
        set_cell_borders(c, bottom=TRIPLE_MID)

    if widths:
        _sync_grid(t, widths)
        for r_ in t.rows:
            cells_u = []
            seen = set()
            for c in r_.cells:
                if id(c._tc) not in seen:
                    seen.add(id(c._tc))
                    cells_u.append(c)
            for i, c in enumerate(cells_u):
                if i < len(widths):
                    c.width = Cm(widths[i])
    if keep_together:
        keep_table_together(t)
    return t


# ---------- 引号修复（只对纯文本 md 用！） ----------
def fix_quotes_cn(text):
    """把直引号 " 按行内出现顺序成对转成中文弯引号。
    ⚠️ 禁止用于 Python 源码文件（会破坏字符串定界符）。"""
    out = []
    for line in text.split('\n'):
        buf, open_q = [], True
        for ch in line:
            if ch == '"':
                buf.append('\u201c' if open_q else '\u201d')
                open_q = not open_q
            else:
                buf.append(ch)
        out.append(''.join(buf))
    return '\n'.join(out)


# ---------- 段落 ----------
def para(doc, text='', cn=SONG, size=12, bold=False, align=None,
         indent=0, sb=0, sa=3, line=1.5):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(sb); pf.space_after = Pt(sa); pf.line_spacing = line
    if indent:
        pf.first_line_indent = Pt(size * indent)
    if text:
        set_font(p.add_run(text), cn, TNR, size, bold)
    return p
