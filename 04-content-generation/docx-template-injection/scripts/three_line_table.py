# -*- coding: utf-8 -*-
"""
把附录三的两个表（表1/表2）改为学术三线表：
  顶线 1.5pt │ 表头下线 0.75pt │ 底线 1.5pt │ 无竖线、无其他横线
只动边框，不改任何内容。直接在用户的文档上原地修改。
"""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

P = r"C:\Users\13662\Desktop\新媒体营销-第一次作业.docx"
doc = Document(P)


def cell_border(cell, edge, sz):
    """给单元格某条边设置边框。edge: top/bottom/left/right。sz 单位=1/8pt"""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    el = borders.find(qn("w:" + edge))
    if el is None:
        el = OxmlElement("w:" + edge)
        borders.append(el)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:color"), "000000")


def cell_no_border(cell, edge):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    el = borders.find(qn("w:" + edge))
    if el is None:
        el = OxmlElement("w:" + edge)
        borders.append(el)
    el.set(qn("w:val"), "nil")


def make_three_line(table):
    n = len(table.rows)
    # 1) 先在表级把所有边框设为 none，清掉 Table Grid 的全框线
    tblPr = table._tbl.tblPr
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        borders.append(el)
    tblPr.append(borders)

    # 2) 单元格级：顶线 + 表头下线 + 底线
    for ci in range(len(table.rows[0].cells)):
        # 表头行：顶线粗 + 下线细
        cell_border(table.cell(0, ci), "top", 12)       # 1.5pt
        cell_border(table.cell(0, ci), "bottom", 6)     # 0.75pt
        # 末行：底线粗
        cell_border(table.cell(n - 1, ci), "bottom", 12)  # 1.5pt
        # 清掉左右竖线
        cell_no_border(table.cell(0, ci), "left")
        cell_no_border(table.cell(0, ci), "right")
        cell_no_border(table.cell(n - 1, ci), "left")
        cell_no_border(table.cell(n - 1, ci), "right")

    for ri in range(1, n):
        for ci in range(len(table.rows[ri].cells)):
            cell = table.cell(ri, ci)
            for edge in ("left", "right", "top"):
                if ri != 0:
                    cell_no_border(cell, edge)
            if ri != n - 1:
                cell_no_border(cell, "bottom")


# ---------- 定位附录三的两个表（它们嵌套在正文单元格内，需递归查找） ----------
def iter_nested_tables(parent_table):
    for row in parent_table.rows:
        for cell in row.cells:
            for t in cell.tables:
                yield t


targets = []
seen = set()
for tbl in doc.tables:
    for t in iter_nested_tables(tbl):
        if id(t._tbl) in seen:
            continue
        seen.add(id(t._tbl))
        try:
            head = [c.text.strip() for c in t.rows[0].cells]
        except Exception:
            continue
        if head[:3] == ["序号", "内容标题", "内容类型"]:
            targets.append((t, "表1"))
        elif head[:2] == ["内容类型", "条数"]:
            targets.append((t, "表2"))

print("找到嵌套目标表格:", [(name, "%d行" % len(t.rows)) for t, name in targets])
assert len(targets) == 2, "目标表格定位失败"

for tbl, name in targets:
    make_three_line(tbl)
    print("  ✔ %s 已改为三线表" % name)

doc.save(P)
print("\n已保存（原地）:", P)
