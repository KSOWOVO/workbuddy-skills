# -*- coding: utf-8 -*-
"""
作业正文注入 v4
- 支持中文图号（图一~图八）
- 支持 @@IMG:文件|图注@@ 案例图标记
- 支持 markdown 表格 → docx 表格（用于附录数据表）
- 支持 ``` 代码块（等宽字体，用于附录爬虫代码）
- 图片直接创建在目标文档的单元格上（关键：保证 rels 正确，图片能显示）
"""
import os, re, shutil, zipfile
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TPL = r"C:\Users\13662\Desktop\新媒体营销-第一次作业(1).docx"
OUTDIR = os.path.join(BASE, "out")
FIGS = os.path.join(BASE, "figs3")
ASSETS = os.path.join(BASE, "assets", "final")
OUT = os.path.join(OUTDIR, "新媒体营销-第一次作业-伍凯森组_终版.docx")
MD = open(os.path.join(OUTDIR, "正文.md"), encoding="utf-8").read()

IMG_W = 5.4

FIGMAP = {
    "图二": ("figB_flow.png",        "图二  本报告的数据采集与处理流程（本组自绘）"),
    "图三": ("figA_pv.png",          "图三  游戏科学官方PV在哔哩哔哩的传播量级（本组采集）"),
    "图四": ("figC_treemap.png",     "图四  内容生态结构：面积表示播放量，颜色区分内容类型（本组采集）"),
    "图五": ("figF_top10.png",       "图五  播放量前十的内容排行（本组采集）"),
    "图六": ("figD_interaction.png", "图六  《黑神话：悟空》13分钟实机演示的互动结构（本组采集）"),
    "图七": ("figE_zhongkui.png",    "图七  《黑神话：钟馗》的传播走势（本组采集）"),
}
# 图一、图八为案例实拍/官方海报，走 @@IMG: 标记


def set_font(run, name="宋体", size=10.5, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rPr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rf.set(qn(a), name)


def sp(p, multiple=1.5, indent=2, before=0, after=3):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = multiple
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if indent:
        pPr = p._element.get_or_add_pPr()
        ind = pPr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind"); pPr.append(ind)
        ind.set(qn("w:firstLineChars"), str(int(indent * 100)))
        ind.set(qn("w:firstLine"), "0")


def body(cell, text):
    p = cell.add_paragraph()
    sp(p, 1.5, 2)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_font(p.add_run(text), "宋体", 10.5)
    return p


def h2(cell, text):
    p = cell.add_paragraph()
    sp(p, 1.5, 0, 10, 5)
    set_font(p.add_run(text), "宋体", 12, True)
    return p


def h3(cell, text):
    p = cell.add_paragraph()
    sp(p, 1.5, 0, 7, 3)
    set_font(p.add_run(text), "宋体", 10.5, True)
    return p


def figure(cell, img, caption):
    p = cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sp(p, 1.0, 0, 8, 2)
    p.add_run().add_picture(img, width=Inches(IMG_W))
    c = cell.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sp(c, 1.0, 0, 0, 10)
    set_font(c.add_run(caption), "宋体", 9)


def code_block(cell, lines):
    for j, ln in enumerate(lines):
        p = cell.add_paragraph()
        sp(p, 1.15, 0, 4 if j == 0 else 0, 4 if j == len(lines) - 1 else 0)
        set_font(p.add_run(ln if ln.strip() else " "), "Consolas", 8.5)


def md_table(cell, rows):
    t = cell.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci)
            for p in list(c.paragraphs):
                p._element.getparent().remove(p._element)
            p = c.add_paragraph()
            sp(p, 1.15, 0, 2, 2)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run(val), "宋体", 9, bold=(ri == 0))
    return t


# ============================ 打开模板 ============================
shutil.copy(TPL, OUT)
doc = Document(OUT)

target = None
for tbl in doc.tables:
    for row in tbl.rows:
        for ci, cell in enumerate(row.cells):
            if cell.text.strip() == "正文" and ci + 1 < len(row.cells):
                target = row.cells[ci + 1]; break
        if target is not None: break
    if target is not None: break
assert target is not None, "未找到正文单元格"

for p in list(target.paragraphs):
    p._element.getparent().remove(p._element)

# ============================ 解析并写入 ============================
lines = MD.split("\n")
i = 0
used_fig, n_case = set(), 0
while i < len(lines):
    s = lines[i].rstrip()
    if not s.strip() or s.startswith("> ") or s.startswith("# "):
        i += 1; continue

    m = re.match(r"^@@IMG:([^|]+)\|([^@]+)@@$", s.strip())
    if m:
        fn, cap = m.group(1).strip(), m.group(2).strip()
        path = os.path.join(ASSETS, fn + ".jpg")
        if os.path.exists(path):
            figure(target, path, cap); n_case += 1
        else:
            print("   [案例图缺失]", path)
        i += 1; continue

    if s.strip().startswith("```"):
        buf = []
        i += 1
        while i < len(lines) and not lines[i].strip().startswith("```"):
            buf.append(lines[i].rstrip()); i += 1
        i += 1
        code_block(target, buf)
        continue

    if s.strip().startswith("|"):
        block = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            block.append(lines[i].strip()); i += 1
        rows = []
        for b in block:
            if re.match(r"^\|[\s:\-|]+\|$", b):
                continue
            rows.append([c.strip() for c in b.strip("|").split("|")])
        ncol = max(len(r) for r in rows)
        rows = [r + [""] * (ncol - len(r)) for r in rows]
        md_table(target, rows)
        target.add_paragraph()
        continue

    if s.startswith("### "):
        h3(target, s[4:].strip()); i += 1; continue
    if s.startswith("## "):
        h2(target, s[3:].strip()); i += 1; continue

    if re.match(r"^\[\d+\]", s.strip()):
        p = target.add_paragraph()
        sp(p, 1.4, 0, 0, 2)
        set_font(p.add_run(s.strip()), "宋体", 9)
        i += 1; continue

    body(target, s)
    for k in re.findall(r"图[一二三四五六七八九十]+", s):
        if k in FIGMAP and k not in used_fig:
            f, cap = FIGMAP[k]
            figure(target, os.path.join(FIGS, f), cap)
            used_fig.add(k)
    i += 1

missing = [k for k in FIGMAP if k not in used_fig]
if missing:
    print("⚠ 未在正文引用：", missing)

doc.save(OUT)
print("数据图 %d 张 | 案例图 %d 张" % (len(used_fig), n_case))
print("已生成:", OUT)

# ============================ 复核 ============================
d2 = Document(OUT)
for tbl in d2.tables:
    for row in tbl.rows:
        for ci, cell in enumerate(row.cells):
            if cell.text.strip() == "正文" and ci + 1 < len(row.cells):
                tc = row.cells[ci + 1]._tc
                txt = "".join(t.text or "" for t in tc.findall(".//" + qn("w:t")))
                print("正文字数:", len(re.findall(r"[\u4e00-\u9fff]", txt)))
                print("图片对象:", len(tc.findall(".//" + qn("w:drawing"))))
                print("内嵌表格:", len(tc.findall(qn("w:tbl"))))

z = zipfile.ZipFile(OUT)
media = [n for n in z.namelist() if n.startswith("word/media/")]
rels = z.read("word/_rels/document.xml.rels").decode("utf-8")
print("word/media:", len(media), "| 图片关系:", rels.count("relationships/image"))
print("文件大小: %.1f MB" % (os.path.getsize(OUT) / 1024 / 1024))
