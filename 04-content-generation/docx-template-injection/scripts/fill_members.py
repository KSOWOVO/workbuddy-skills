# -*- coding: utf-8 -*-
"""填写封面、同组人员、小组分工表"""
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(BASE, "out", "新媒体营销-第一次作业-伍凯森组.docx")

MEMBERS = ["伍凯森", "陈思全", "杨婷", "蒲俊竹", "徐惠清", "张佳颖", "李乐菲"]
DUTIES = [
    "选题策划、平台数据采集与核验、理论框架搭建、报告统稿与格式排版",
    "案例背景资料收集、企业与产品基本情况梳理",
    "目标用户画像分析、平台选择策略论证",
    "内容创意策略分析、传播节奏与节点运营梳理",
    "用户互动数据整理、弹幕与评论区互动特征分析",
    "营销效果评估、数据图表制作与可视化呈现",
    "存在问题诊断、优化建议撰写、文字校对",
]


def set_font(run, name="宋体", size=10.5, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rPr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rf.set(qn(a), name)


def fill_cell(cell, text, size=10.5, bold=False, align="left"):
    """清空并填充单元格，统一格式"""
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)
    p = cell.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_before = Pt(0); pf.space_after = Pt(0)
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT,
                   "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    r = p.add_run(text)
    set_font(r, "宋体", size, bold)


doc = Document(OUT)

# ---------- TABLE 0：封面 ----------
t0 = doc.tables[0]
fill_cell(t0.rows[4].cells[1], "伍凯森  等", size=15, bold=True, align="center")

# ---------- TABLE 1：同组人员 ----------
t1 = doc.tables[1]
fill_cell(t1.rows[2].cells[3], "\n".join(MEMBERS), size=10.5)

# ---------- TABLE 2：分工表 ----------
t2 = doc.tables[2]
# 表头第0行保持不动
for i, (name, duty) in enumerate(zip(MEMBERS, DUTIES)):
    row = t2.rows[i + 1]
    fill_cell(row.cells[0], "", size=10.5, align="center")     # 学号留空（用户后补）
    fill_cell(row.cells[1], name, size=10.5, align="center")
    fill_cell(row.cells[2], "", size=10.5, align="center")     # 成绩由教师填
    fill_cell(row.cells[3], duty, size=10.5)

doc.save(OUT)
print("封面/同组人员/分工表 已填写完成")
print("成员:", " ".join(MEMBERS))

# 复核
d2 = Document(OUT)
print("\n--- 封面姓名:", d2.tables[0].rows[4].cells[1].text)
print("--- 同组人员:", repr(d2.tables[1].rows[2].cells[3].text))
print("--- 分工表 ---")
for r in d2.tables[2].rows[1:]:
    print("  [%s] %s" % (r.cells[1].text, r.cells[3].text[:40]))
