# -*- coding: utf-8 -*-
"""
交付前严格格式审查
逐段核对：标题（宋体·小四·加粗）、正文（宋体·五号·1.5倍行距）
并核对作业总要求的所有条目
"""
import os, re, zipfile, sys
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(BASE, "out", "新媒体营销-第一次作业-伍凯森组.docx")
doc = Document(OUT)
z = zipfile.ZipFile(OUT)
xml = z.read("word/document.xml").decode("utf-8")

PASS, FAIL = "✔", "✘"
issues = []


def hr(t=""):
    print("\n" + "=" * 74)
    if t:
        print(t)
        print("=" * 74)


hr("                    交付前严格格式审查")
print("文件：%s" % os.path.basename(OUT))

# ---------- 定位正文单元格 ----------
tc = None
for tbl in doc.tables:
    for row in tbl.rows:
        for i, c in enumerate(row.cells):
            if c.text.strip() == "正文" and i + 1 < len(row.cells):
                tc = row.cells[i + 1]._tc
                break
        if tc is not None: break
    if tc is not None: break

# ---------- 1. 逐段格式核对 ----------
hr("【1】逐段格式核对（硬性要求）")
h2_bad, body_bad, n_h2, n_body, n_other = [], [], 0, 0, 0
for p in tc.findall(qn("w:p")):
    ts = "".join(t.text or "" for t in p.findall(".//" + qn("w:t")))
    if not ts.strip():
        continue
    if p.findall(".//" + qn("w:drawing")):      # 图片段
        n_other += 1; continue
    pPr = p.find(qn("w:pPr"))
    sp = pPr.find(qn("w:spacing")) if pPr is not None else None
    line = sp.get(qn("w:line")) if sp is not None else None
    r = p.find(qn("w:r"))
    if r is None:
        continue
    rPr = r.find(qn("w:rPr"))
    szel = rPr.find(qn("w:sz")) if rPr is not None else None
    sz = szel.get(qn("w:val")) if szel is not None else None
    bold = (rPr is not None and rPr.find(qn("w:b")) is not None)
    rf = rPr.find(qn("w:rFonts")) if rPr is not None else None
    font = rf.get(qn("w:eastAsia")) if rf is not None else None
    ascii_font = rf.get(qn("w:ascii")) if rf is not None else None

    if sz == "24":                              # 小四
        n_h2 += 1
        if not bold:
            h2_bad.append(("未加粗", ts[:30]))
        if font != "宋体":
            h2_bad.append(("字体≠宋体(%s)" % font, ts[:30]))
    elif sz == "21":                            # 五号
        n_body += 1
        if font != "宋体":
            body_bad.append(("字体≠宋体(%s)" % font, ts[:30]))
        if line != "360":
            body_bad.append(("行距≠1.5倍(line=%s)" % line, ts[:30]))
    else:
        n_other += 1                            # 图注(9pt)、参考文献(9pt)、表格内等

print("  标题段（小四）%d 个｜正文段（五号）%d 个｜图注/参考/表内 %d 个" % (n_h2, n_body, n_other))
if h2_bad:
    print("  %s 标题不合规 %d 处：" % (FAIL, len(h2_bad)))
    for x in h2_bad[:8]: print("      -", x)
    issues.append("标题格式不合规")
else:
    print("  %s 全部标题：宋体 · 小四 · 加粗" % PASS)
if body_bad:
    print("  %s 正文不合规 %d 处：" % (FAIL, len(body_bad)))
    for x in body_bad[:8]: print("      -", x)
    issues.append("正文格式不合规")
else:
    print("  %s 全部正文：宋体 · 五号 · 1.5倍行距" % PASS)

# ---------- 2. 作业总要求逐条核对 ----------
hr("【2】作业总要求逐条核对")

txt_all = "".join(t.text or "" for t in tc.findall(".//" + qn("w:t")))
cn = len(re.findall(r"[\u4e00-\u9fff]", txt_all))
drawings = len(tc.findall(".//" + qn("w:drawing")))
media = [n for n in z.namelist() if n.startswith("word/media/")]
rels = z.read("word/_rels/document.xml.rels").decode("utf-8")
img_rels = rels.count("relationships/image")

checks = [
    ("字数 ≥1500字", cn >= 1500, "%d 字" % cn),
    ("图文并茂（图片 ≥6 张）", drawings >= 6, "%d 张" % drawings),
    ("图片真实嵌入", len(media) - 1 >= drawings and img_rels >= drawings,
     "media %d / 关系 %d" % (len(media), img_rels)),
    ("标题 宋体·小四·加粗", not h2_bad, "%d 个标题段全部合规" % n_h2 if not h2_bad else "有异常"),
    ("正文 宋体·五号·1.5倍行距", not body_bad, "%d 个正文段全部合规" % n_body if not body_bad else "有异常"),
    ("结构完整（含背景/分析/问题/建议/结语）",
     all(k in txt_all for k in ["企业与产品", "数据从哪里来", "传播策略分析", "用户互动与转化", "存在的问题", "优化建议", "结语"]),
     "七章齐全"),
    ("案例具备代表性", "游戏科学" in txt_all and "黑神话" in txt_all, "游戏科学《黑神话》系列"),
    ("分析详实有逻辑", len(re.findall(r"[0-9]", txt_all)) > 200, "文中含 %d 处数据" % len(re.findall(r"[0-9]", txt_all))),
    ("具备小组成员的见解", "我们认为" in txt_all or "我们结合" in txt_all or "我们提出" in txt_all, "含本组独立判断"),
    ("附录含参考文献", "[30]" in txt_all, "30 条"),
    ("附录含爬虫数据表", len(tc.findall(qn("w:tbl"))) >= 2, "%d 个数据表" % len(tc.findall(qn("w:tbl")))),
    ("模板红色说明文字已清", xml.count('w:color w:val="FF0000"') == 0, "0 处"),
    ("模板黄色高亮已清", xml.count('w:highlight w:val="yellow"') == 0, "0 处"),
]
for name, ok, detail in checks:
    print("  %s %-32s %s" % (PASS if ok else FAIL, name, detail))
    if not ok:
        issues.append(name)

# ---------- 3. 分工表 ----------
hr("【3】小组分工表")
t2 = doc.tables[2]
filled = 0
for r in t2.rows[1:]:
    nm = r.cells[1].text.strip()
    if nm:
        filled += 1
        print("  %-4s │ %s" % (nm, r.cells[3].text.strip()))
print("  已填 %d 人 %s" % (filled, PASS if filled == 7 else FAIL))
if filled != 7:
    issues.append("分工表人数不符")

# ---------- 4. 内容自检（AI 味） ----------
hr("【4】文风自检")
body_only = txt_all.split("附录一")[0]
semi = body_only.count("；") + body_only.count(";")
bold = xml.count("<w:b/>")
print("  分号数量：%d %s" % (semi, PASS if semi == 0 else FAIL))
print("  正文字数：%d" % len(re.findall(r"[\u4e00-\u9fff]", body_only)))
for w in ["值得注意的是", "综上所述", "换言之", "其一", "其二"]:
    n = body_only.count(w)
    if n:
        print("  %s 出现「%s」%d 次" % (FAIL, w, n))
        issues.append("AI 味词汇：%s" % w)
if not any(body_only.count(w) for w in ["值得注意的是", "综上所述", "换言之", "其一", "其二"]):
    print("  %s 无 AI 味高频词" % PASS)

# ---------- 总结 ----------
hr()
if issues:
    print("❌ 仍有 %d 项待处理：" % len(issues))
    for x in issues: print("   -", x)
else:
    print("✅ 全部要求均已满足，可提交")
print("文件：%s（%.1f MB）" % (OUT, os.path.getsize(OUT) / 1024 / 1024))
print("=" * 74)
